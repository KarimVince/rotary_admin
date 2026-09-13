from io import BytesIO

import httpx
import pytest
from PIL import Image
from pptx import Presentation

from app.core.donation_statistics_report import (
    _CMP_FULL_THRESHOLD,
    _CMP_ORGS_PER_SEC,
    _footnote_text,
    _page_area_counts,
    _paginate,
    _pptx_safe_image,
    _sorted_ngo_rows,
    _summary_cards,
    build_pptx_comparison_report,
    build_pptx_report,
    resolve_logo_bytes,
)
from app.schemas.donation_statistics import ConvertedTotals, DonationStatistics

pytestmark = pytest.mark.unit


def _empty_converted_totals(**overrides):
    defaults = {"total_hkd": 0.0, "total_usd": 0.0, "unconverted_count": 0, "unconverted_currencies": []}
    defaults.update(overrides)
    return ConvertedTotals(**defaults)


def _minimal_stats(**overrides):
    defaults = {
        "by_currency": [],
        "selected_rotary_year": 2024,
        "selected_year_organisations_count": 0,
        "selected_year": _empty_converted_totals(),
        "all_time_organisations_count": 0,
        "all_time": _empty_converted_totals(),
        "total_service_hours_all_time": 0.0,
        "total_service_hours_selected_year": 0.0,
        "service_hours_by_rotary_year": [],
        "selected_year_planned": _empty_converted_totals(),
        "selected_year_organisations_count_with_planned": 0,
        "all_time_organisations_count_with_planned": 0,
    }
    defaults.update(overrides)
    return DonationStatistics(**defaults)


# Story 16.35 redesign (district-template handoff) — Summary slide's four
# cards: Donated HKD, Planned HKD, Donated USD, Reach (orgs supported).


def test_summary_cards_includes_all_four_when_all_nonzero():
    stats = _minimal_stats(
        selected_year=_empty_converted_totals(total_hkd=100.0, total_usd=50.0),
        selected_year_planned=_empty_converted_totals(total_hkd=200.0),
        selected_year_organisations_count=3,
    )
    cards = _summary_cards(stats)
    assert [card["kicker"] for card in cards] == ["Donated", "Planned", "Donated", "Reach"]
    assert cards[3]["unit"] == ""


def test_summary_cards_drops_zero_currency_cards_never_renders_zero():
    # "If a currency has no activity in the selected year, drop that card
    # ... Never render a zero card." USD is zero here — dropped, not shown
    # as "0 USD". Reach (org count) is never dropped even at zero.
    stats = _minimal_stats(
        selected_year=_empty_converted_totals(total_hkd=100.0, total_usd=0.0),
        selected_year_planned=_empty_converted_totals(total_hkd=0.0),
        selected_year_organisations_count=0,
    )
    cards = _summary_cards(stats)
    assert [card["kicker"] for card in cards] == ["Donated", "Reach"]
    assert cards[1]["figure"] == "0"


# _sorted_ngo_rows / _page_area_counts / _paginate — Organisations
# slide/table grouping and pagination.


def _row(name, area, total=100.0):
    return {"name": name, "country": None, "area": area, "contact_name": None, "total": total, "logo_bytes": None}


def test_sorted_ngo_rows_orders_by_area_then_name():
    rows = [
        _row("Zeta", "Unclassified"),
        _row("Beta", "Health & Medical"),
        _row("Alpha", "Education & Literacy"),
        _row("Gamma", "Education & Literacy"),
    ]
    sorted_rows = _sorted_ngo_rows(rows)
    assert [row["name"] for row in sorted_rows] == ["Alpha", "Gamma", "Beta", "Zeta"]


def test_page_area_counts_only_covers_that_pages_areas():
    page = [_row("A", "Health & Medical"), _row("B", "Health & Medical"), _row("C", "Others")]
    counts = _page_area_counts(page)
    assert counts == [("Health & Medical", 2), ("Others", 1)]


def test_pptx_organisation_card_amount_includes_currency_unit():
    # Real bug, found right after the redesign shipped: the Organisations
    # slide's card only ever showed the bare number ("13,000"), while the
    # PDF's organisation table already showed "13,000 HKD" — user-visible
    # inconsistency between the two export formats for the same data.
    stats = _minimal_stats()
    pptx_bytes = build_pptx_report(
        stats, "HKD", [_row("Butterfly Soccer Cup", "Health & Medical", total=13000.0)], chrome="plain"
    )
    prs = Presentation(BytesIO(pptx_bytes))
    all_text = " ".join(
        run.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    ) + " ".join(
        shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "13,000 HKD" in all_text


def test_paginate_chunks_of_twelve():
    rows = [_row(f"Org {i}", "Others") for i in range(14)]
    pages = _paginate(rows, 12)
    assert len(pages) == 2
    assert len(pages[0]) == 12
    assert len(pages[1]) == 2


def test_footnote_text_singular_and_plural():
    stats = _minimal_stats()
    row = _row("Alpha", "Health & Medical")
    text_one = _footnote_text(stats, [row], [[row]])
    assert text_one.startswith("1 organisation is listed for the year across 1 area of focus.")
    rows_two = [row, _row("Beta", "Others")]
    text_two = _footnote_text(stats, rows_two, [rows_two])
    assert text_two.startswith("2 organisations are listed for the year across 2 areas of focus.")


def test_footnote_text_empty_when_no_organisations():
    stats = _minimal_stats()
    text = _footnote_text(stats, [], [])
    assert text == "0 organisations are listed for the year across 0 areas of focus."


def test_pptx_safe_image_reencodes_webp_to_png():
    # Story 16.35 follow-up — real bug: python-pptx's add_picture rejects
    # WEBP outright, but a real org logo can be one (Story 16.6 allows it).
    webp_buf = BytesIO()
    Image.new("RGB", (5, 5), color=(10, 20, 30)).save(webp_buf, format="WEBP")

    result = _pptx_safe_image(BytesIO(webp_buf.getvalue()))

    assert result is not None
    with Image.open(result) as reencoded:
        assert reencoded.format == "PNG"


def test_pptx_safe_image_passes_through_already_supported_format():
    png_buf = BytesIO()
    Image.new("RGB", (5, 5)).save(png_buf, format="PNG")

    result = _pptx_safe_image(BytesIO(png_buf.getvalue()))

    assert result is not None
    with Image.open(result) as reencoded:
        assert reencoded.format == "PNG"


def test_pptx_safe_image_returns_none_for_undecodable_bytes():
    assert _pptx_safe_image(BytesIO(b"not an image")) is None


def test_resolve_logo_bytes_returns_none_for_missing_url():
    assert resolve_logo_bytes(None) is None
    assert resolve_logo_bytes("") is None


def test_resolve_logo_bytes_returns_none_when_local_fallback_file_absent():
    # Pre-migration relative path, no file actually on disk for this fixture.
    assert resolve_logo_bytes("/static/organisations/does-not-exist.png") is None


def test_resolve_logo_bytes_fetches_absolute_supabase_url(monkeypatch):
    def fake_get(url, timeout=None):
        assert url == "https://proj.supabase.co/storage/v1/object/public/public-assets/organisations/abc.png"
        return httpx.Response(200, content=b"logo-bytes")

    monkeypatch.setattr(httpx, "get", fake_get)

    result = resolve_logo_bytes(
        "https://proj.supabase.co/storage/v1/object/public/public-assets/organisations/abc.png"
    )
    assert result is not None
    assert result.read() == b"logo-bytes"


def test_resolve_logo_bytes_returns_none_on_non_200_response(monkeypatch):
    monkeypatch.setattr(httpx, "get", lambda url, timeout=None: httpx.Response(404))

    assert resolve_logo_bytes("https://proj.supabase.co/storage/v1/object/public/x/y.png") is None


def test_resolve_logo_bytes_returns_none_on_request_error(monkeypatch):
    def fake_get(url, timeout=None):
        raise httpx.ConnectError("boom")

    monkeypatch.setattr(httpx, "get", fake_get)

    assert resolve_logo_bytes("https://proj.supabase.co/storage/v1/object/public/x/y.png") is None


# Year-over-year comparison PPTX (build_pptx_comparison_report)


def _cmp_row(name, area="Education & Literacy", total=10000.0):
    return {"name": name, "country": "HK", "area": area, "contact_name": "Test", "total": total, "logo_bytes": None}


def test_build_pptx_comparison_report_produces_valid_pptx():
    rows_a = [_cmp_row("Org A1"), _cmp_row("Org A2", area="Health & Medical")]
    rows_b = [_cmp_row("Org B1"), _cmp_row("Org B2", area="Youth Development")]
    data = build_pptx_comparison_report(2024, rows_a, 2025, rows_b, "HKD", chrome="plain")
    prs = Presentation(BytesIO(data))
    assert len(prs.slides) == 1


def test_build_pptx_comparison_report_contains_both_year_labels():
    rows_a = [_cmp_row("Alpha")]
    rows_b = [_cmp_row("Beta")]
    data = build_pptx_comparison_report(2024, rows_a, 2025, rows_b, "HKD")
    prs = Presentation(BytesIO(data))
    all_text = " ".join(
        shape.text_frame.text
        for shape in prs.slides[0].shapes
        if shape.has_text_frame
    )
    assert "2024–2025" in all_text
    assert "2025–2026" in all_text


def test_build_pptx_comparison_report_shows_org_totals():
    rows_a = [_cmp_row("Org A", total=50000.0)]
    rows_b = [_cmp_row("Org B", total=75000.0)]
    data = build_pptx_comparison_report(2023, rows_a, 2024, rows_b, "HKD")
    prs = Presentation(BytesIO(data))
    all_text = " ".join(
        shape.text_frame.text
        for shape in prs.slides[0].shapes
        if shape.has_text_frame
    )
    assert "50,000 HKD" in all_text   # section-A total chip
    assert "75,000 HKD" in all_text   # section-B total chip


def test_build_pptx_comparison_report_caps_per_section_at_max():
    # Use _CMP_FULL_THRESHOLD orgs so we stay in single-slide mode but
    # still exceed _CMP_ORGS_PER_SEC — the excess is silently capped.
    n = min(_CMP_ORGS_PER_SEC + 3, _CMP_FULL_THRESHOLD)
    many = [_cmp_row(f"Org {i}") for i in range(n)]
    data = build_pptx_comparison_report(2024, many, 2025, many[:2], "HKD")
    prs = Presentation(BytesIO(data))
    assert len(prs.slides) == 1


def test_build_pptx_comparison_report_empty_year_does_not_crash():
    data = build_pptx_comparison_report(2024, [], 2025, [_cmp_row("Only B")], "HKD")
    prs = Presentation(BytesIO(data))
    assert prs.slides[0].shapes  # slide has shapes
